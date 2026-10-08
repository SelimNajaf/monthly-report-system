"""Read .docx reports into a structured, image-free text model.

Parses a Word document with ``python-docx`` and exposes its tables as plain 2D
grids of cell text. The real inspection reports carry *all* their data in tables
(header metadata block, weld-data table, signature block) and zero body prose, so
tables are the only payload we need. Images are never read into the model, which
satisfies the "ignore all images" requirement for free.

I/O contract
------------
read_tables(path) -> list[list[list[str]]]
    A list of tables; each table is a list of rows; each row is a list of cell
    strings (cell text = its paragraphs joined by a newline, stripped).
read_document(path) -> (tables, has_text)
    The same tables, plus a flag saying whether the document body holds any
    readable text at all (False = a scanned image pasted into Word).
read_many(paths) -> (cache, errors, textless)
    Bulk parallel read; ``textless`` lists the files that parsed cleanly but
    carry no readable text.
"""

from __future__ import annotations

import os
import re
import sys
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .console import progress_printer

#: Stage label for the live progress counter. Both halves of STAGE 1 report under
#: it — the parse-cache hash pass (parse_cache.read_many_cached) and the pool that
#: parses whatever the cache missed — because to the operator they are one stage.
#: They keep separate denominators, which is honest: the second pass only covers
#: the misses, and pretending otherwise would show a bar that leaps then crawls.
STAGE_PARSE = "Stage 1: Parsing"


def parse_filename_metadata(path: str | Path) -> dict[str, str | None]:
    """Extract official Document Name, Report № and Standard from the file *name*.

    Inspectors sometimes mistype the in-document standard field, so the filename
    is the authoritative source. Layout: ``<name> (<job>) (<standard>)_en.docx``.

      KF-MT-26-003 (2230) (ISO-23278)_en.docx
        -> report_no = "KF-MT-26-003",
           document_name = "KF-MT-26-003 (2230)", standard = "ISO-23278"
      B-MT-26-001-lan.docx                      (LAN no-clobber twin)
        -> report_no = "B-MT-26-001"
      U-UT-26-010 R1 (Platform)_en_en.docx      (only one paren group)
        -> report_no = "U-UT-26-010 R1",
           document_name = "U-UT-26-010 R1 (Platform)", standard = None (use body)

    ``report_no`` is the bare name before the first parenthesis group (the
    authoritative report number used for the filename↔internal cross-check).

    Robust: any layout that does not match still yields a sensible
    ``document_name`` and ``standard=None`` — never raises.
    """
    stem = Path(path).stem
    # A ``-lan`` / ``-lanN`` tail is the LAN sync's no-clobber twin marker, never
    # part of a report number. With a paren group it already falls outside
    # ``name``; without one (``B-MT-26-001-lan``) it would leak into the number.
    stem = re.sub(r"-lan\d*$", "", stem, flags=re.IGNORECASE).strip()
    stem = re.sub(r"(_en)+$", "", stem, flags=re.IGNORECASE).strip()
    name = re.split(r"\(", stem, maxsplit=1)[0].strip()
    parens = re.findall(r"\(([^)]*)\)", stem)
    parens = [p.strip() for p in parens if p.strip()]

    document_name = f"{name} ({parens[0]})" if parens else (name or stem)
    standard = parens[1] if len(parens) >= 2 else None
    return {"report_no": name or stem, "document_name": document_name, "standard": standard}


def _cell_text(cell) -> str:
    """Concatenate a table cell's paragraph text (images/drawings contribute nothing)."""
    parts = [p.text for p in cell.paragraphs]
    return "\n".join(t for t in parts).strip()


def has_body_text(doc) -> bool:
    """True when the document *body* contains at least one non-blank text run.

    Scans the body XML for ``<w:t>`` rather than walking ``doc.paragraphs``,
    which is deliberate on both ends of the scope:

    * **Wider:** ``<w:t>`` also catches text inside tables and text boxes
      (``w:txbxContent``). PDF-to-Word converters routinely emit text boxes, and
      those are invisible to ``doc.paragraphs`` — a real converted report would
      otherwise be misjudged as a scan.
    * **Narrower:** headers and footers live in separate OOXML parts, so they are
      excluded for free. A pasted scan often still sits under a Word header
      carrying the company name; counting that text would mask every scan.
    """
    return any((t.text or "").strip() for t in doc.element.body.iter(qn("w:t")))


def read_document(path: str | Path) -> tuple[list[list[list[str]]], bool]:
    """Parse once; return ``(tables, has_text)``.

    Single place where a .docx is opened, so the text probe costs no extra I/O.
    """
    doc = Document(str(path))
    tables: list[list[list[str]]] = []
    for table in doc.tables:
        grid: list[list[str]] = []
        for row in table.rows:
            # row.cells expands merged cells by repetition; collapse consecutive
            # duplicates so a horizontally-merged label is not counted many times.
            texts: list[str] = []
            prev_id = None
            for cell in row.cells:
                cid = id(cell._tc)
                if cid == prev_id:
                    continue
                prev_id = cid
                texts.append(_cell_text(cell))
            grid.append(texts)
        tables.append(grid)
    return tables, has_body_text(doc)


def read_tables(path: str | Path) -> list[list[list[str]]]:
    """Return every table in the document as a 2D grid of cell text."""
    return read_document(path)[0]


# --- parallel bulk reading --------------------------------------------------
#
# The pipeline used to parse every .docx 2–3× (once to discover date/method
# groups, again per compiled group). ``read_many`` parses each file **once**,
# in parallel across CPU cores, and hands back a cache the rest of the pipeline
# reuses. The .docx parse is CPU-bound XML (lxml), so processes give a real
# multi-core speed-up on the 100s of reports in a month.

def _worker_init(src_dir: str) -> None:
    """Ensure the ``monthly_report`` package is importable in a spawned worker.

    On Windows (spawn), a fresh interpreter does not inherit the parent's
    ``sys.path`` edits, so we re-add ``src/`` here before any task runs.
    """
    if src_dir and src_dir not in sys.path:
        sys.path.insert(0, src_dir)


def _read_one(path_str: str):
    """Parse one file; return ``(path_str, tables, has_text, disagreements)`` or
    ``(path_str, error_str, False, [])``.

    Dispatches on the file suffix. A ``.docx`` travels the path it always has —
    the same ``read_document`` above, byte for byte. A ``.pdf`` is read directly
    (ADR-003) by two independent readers, and whatever they disagree about comes
    back on the fourth element for the caller to act on; the tables returned are
    always the primary reader's.

    The PDF imports are function-level on purpose. They pull in PyMuPDF and, via
    ``pdf_crosscheck``, the aggregator — which imports this module. Importing them
    here breaks that cycle and keeps a .docx-only run from loading PyMuPDF at all.

    Never raises — a corrupt/locked file yields its error string so the caller
    can quarantine it without aborting the whole batch.
    """
    try:
        if Path(path_str).suffix.lower() == ".pdf":
            from . import pdf_crosscheck, pdf_reader
            primary, crosscheck, has_text = pdf_reader.read_both(path_str)
            found = pdf_crosscheck.compare(primary, crosscheck, path=path_str)
            return path_str, primary, has_text, found
        tables, has_text = read_document(path_str)
        return path_str, tables, has_text, []
    except Exception as exc:  # unreadable / corrupt / locked
        return path_str, f"{type(exc).__name__}: {exc}", False, []


def read_many(paths, max_workers: int | None = None, progress_every: int = 0):
    """Parse many .docx in parallel.

    Returns ``(cache, errors, textless)`` where ``cache`` is ``{Path: tables}``
    for files that parsed cleanly, ``errors`` is ``{Path: message}`` for those
    that did not, and ``textless`` lists the cleanly-parsed files whose body
    holds no readable text (a scanned image pasted into Word).

    A textless file parsed *fine*, so it stays in ``cache`` and out of
    ``errors``: the two conditions need different advice, and the caller decides
    what to do with each. Falls back to a serial loop if a process pool cannot be
    started, so the run always completes (parallelism is an optimisation, never a
    correctness dependency).

    ``progress_every`` turns on the live ``[Stage 1: Parsing] [i/N]`` counter
    (0 = silent, as every existing caller expects). The index is assigned in this
    parent process as results land, never by a worker, so it stays monotonic —
    the same discipline pre_processor uses for the PDF counter.

    PDF cross-check disagreements (ADR-003) are published on
    ``read_many.last_disagreements`` rather than returned, so the triple every
    existing caller unpacks is unchanged — the same side-channel pattern
    ``parse_cache.read_many_cached.last_stats`` uses.
    """
    paths = [Path(p) for p in paths]
    cache: dict[Path, list] = {}
    errors: dict[Path, str] = {}
    textless: list[Path] = []
    disagreements: list = []
    read_many.last_disagreements = disagreements
    if not paths:
        return cache, errors, textless

    src_dir = str(Path(__file__).resolve().parent.parent)  # .../src
    workers = max_workers or (os.cpu_count() or 1)
    workers = max(1, min(workers, len(paths)))

    emit = progress_printer(STAGE_PARSE, len(paths), progress_every)
    done = 0

    def _record(path_str, result, has_text, found) -> None:
        nonlocal done
        p = Path(path_str)
        done += 1
        if isinstance(result, str):
            emit(done, f"Parse failed: {p.name} ({result})")
            errors[p] = result
            return
        emit(done, f"Parsed report: {p.name}")
        cache[p] = result
        if not has_text:
            textless.append(p)
        disagreements.extend(found)

    try:
        with ProcessPoolExecutor(max_workers=workers,
                                 initializer=_worker_init,
                                 initargs=(src_dir,)) as pool:
            for record in pool.map(_read_one, [str(p) for p in paths]):
                _record(*record)
    except Exception as exc:  # pool unavailable -> serial fallback, still correct
        print(f"  [read_many] parallel read unavailable ({exc}); reading serially")
        cache.clear()
        errors.clear()
        textless.clear()
        disagreements.clear()
        done = 0          # the batch restarts; so does its counter
        for p in paths:
            _record(*_read_one(str(p)))
    return cache, errors, textless


#: Published by the most recent ``read_many`` call: the PDF cross-check
#: disagreements it found (ADR-003). Empty before the first call and after any
#: batch of pure .docx, so a caller may read it unconditionally.
read_many.last_disagreements = []
