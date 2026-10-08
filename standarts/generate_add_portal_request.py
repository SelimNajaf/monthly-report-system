#!/usr/bin/env python3
"""Generate the official "add to E-Keyfiyyet portal" request letter.

Emits BOTH ``add_portal_request.docx`` (python-docx) and ``add_portal_request.pdf``
(ReportLab) from ONE shared dataset + one set of shared Azerbaijani strings, so the two
formats never drift. The 21 rows correspond to ``add_portal.md`` — the standards proposed
for addition to the portal's normative-document registry.

The enriched columns (full title + NDT method) are authored from NDT domain knowledge:
the repository has no standard->title mapping. Rows whose Russian-GOST title could not be
confirmed are marked ``uncertain=True`` and printed at the end for operator review before
the letter is sent.

Azerbaijani Unicode (ə ğ ı İ ö ş ç ü) needs a covering TTF; we use the bundled
``assets/fonts/DejaVuSans.ttf`` / ``DejaVuSans-Bold.ttf``. Font paths are resolved relative
to the repo root (derived from ``__file__``) so the script runs from any working directory,
and it fails fast with a clear message if a TTF is missing — mirroring the renderer contract
in ``assets/fonts/README.md``.

Run:  python standarts/generate_add_portal_request.py
"""

from __future__ import annotations

import sys
from pathlib import Path

# Windows consoles default to a legacy codepage (cp1251) that cannot encode the
# Azerbaijani/Cyrillic characters in our status output; force UTF-8 so prints never crash.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

# --- Paths -----------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent          # .../standarts
REPO_ROOT = SCRIPT_DIR.parent                          # repo root
FONT_REGULAR = REPO_ROOT / "assets" / "fonts" / "DejaVuSans.ttf"
FONT_BOLD = REPO_ROOT / "assets" / "fonts" / "DejaVuSans-Bold.ttf"
OUT_DOCX = SCRIPT_DIR / "add_portal_request.docx"
OUT_PDF = SCRIPT_DIR / "add_portal_request.pdf"

# --- Shared strings --------------------------------------------------------------------

TITLE = (
    "E-Keyfiyyet Portalına Əlavə Edilməsi Təklif Olunan Beynəlxalq və "
    "Dövlət Standartlarının Siyahısı"
)
INTRO = (
    "Qeyri-Dağıdıcı Müayinə (QNM) fəaliyyətimizdə istinad edilən, lakin E-Keyfiyyet "
    "portalının normativ-sənədlər reyestrində hazırda mövcud olmayan aşağıdakı beynəlxalq "
    "və dövlət standartlarının portala əlavə edilməsini təklif edirik. Standartların tam "
    "adları və aid olduqları QNM (NDT) metodları cədvəldə göstərilmişdir."
)
COL_HEADERS = ("№", "Standartın Tam Adı", "Aid Olduğu NDT Metodu / Təsvir")
SIGNATURE_LABELS = ("Ad, Soyad", "Vəzifə", "İmza", "Tarix")

# --- Dataset: (no, code, full_name_az, ndt_method_desc_az, uncertain) -------------------
# QNM = Qeyri-Dağıdıcı Müayinə (NDT). Methods: RT radioqrafik, UT ultrasəs,
# MT maqnit-toz, PT/DP kapillyar (nüfuzedici), VT vizual, PMI material analizi,
# PAUT faza-qəfəsli UT, TOFD diffraksiya, UTM qalınlıq ölçmə.

STANDARDS: list[tuple[int, str, str, str, bool]] = [
    (1,  "API 510",           "Təzyiq qablarının istismar müayinəsi kodeksi",
         "Təzyiq qablarının ümumi müayinəsi (VT, UT, RT, MT, PT)", False),
    (2,  "API 570",           "Boru kəmərlərinin istismar müayinəsi kodeksi",
         "Boru kəmərlərinin ümumi müayinəsi (VT, UT, RT, MT, PT)", False),
    (3,  "API 620",           "İri həcmli qaynaqlı aşağı təzyiqli saxlama çənlərinin "
                              "layihələndirilməsi və tikintisi",
         "Çən qaynaqlarının müayinəsi (RT, UT, VT)", False),
    (4,  "API 650",           "Neft saxlanması üçün qaynaqlı çənlər",
         "Çən qaynaqlarının müayinəsi (RT, UT, VT, MT, PT)", False),
    (5,  "ASME B31.1",        "Enerji boru kəmərləri (Power Piping)",
         "Boru qaynaqlarının müayinəsi (RT, UT, MT, PT, VT)", False),
    (6,  "ASME B31.3",        "Proses boru kəmərləri (Process Piping)",
         "Boru qaynaqlarının müayinəsi (RT, UT, MT, PT, VT)", False),
    (7,  "ASME I-XII",        "ASME Qazan və Təzyiq Qabları Kodeksi (Bölmə I–XII)",
         "Qaynaq və material üçün ümumi qəbul kodeksi (RT, UT, MT, PT, VT)", False),
    (8,  "ASTM E 572",        "Paslanmayan və legirlənmiş poladların rentgen-floresans "
                              "(XRF) spektrometriyası ilə analizi",
         "PMI — material tərkibinin təyini", False),
    (9,  "ASTM E 1476",       "Metalların identifikasiyası, markasının təsdiqi və "
                              "çeşidlənməsi üzrə təlimat",
         "PMI — material identifikasiyası", False),
    (10, "AWS D1.1",          "Konstruksiya Qaynaq Kodeksi — Polad",
         "Konstruksiya qaynaqlarının müayinəsi (VT, RT, UT, MT, PT)", False),
    (11, "ISO 3057",          "QNM — Səth müayinəsi üçün metaloqrafik replika üsulları",
         "Səth/replika müayinəsi", False),
    (12, "ISO 10863",         "Qaynaqların QNM — Ultrasəs müayinəsi — TOFD üsulunun "
                              "tətbiqi",
         "UT (TOFD — diffraksiya vaxtı üsulu)", False),
    (13, "ISO 16809",         "QNM — Ultrasəs ilə qalınlığın ölçülməsi",
         "UTM — ultrasəs qalınlıq ölçmə", False),
    (14, "ISO 17637",         "Qaynaqların QNM — Əridilmiş qaynaq birləşmələrinin vizual "
                              "müayinəsi",
         "VT — vizual müayinə", False),
    (15, "ISO 19285",         "Qaynaqların QNM — Faza-qəfəsli ultrasəs müayinəsi (PAUT) — "
                              "Qəbul səviyyələri",
         "UT (PAUT — faza-qəfəsli)", False),
    (16, "ГОСТ 7512-82",      "QNM. Qaynaq birləşmələri. Radioqrafik üsul",
         "RT — radioqrafik müayinə", False),
    (17, "ГОСТ 18442-80",     "QNM. Kapillyar üsullar. Ümumi tələblər",
         "PT — kapillyar (nüfuzedici) müayinə", False),
    (18, "ГОСТ 26790",        "QNM. Kapillyar üsullar",
         "PT — kapillyar (nüfuzedici) müayinə", True),
    (19, "ГОСТ Р 51780",      "QNM sahəsinə aid dövlət standartı",
         "Dəqiqləşdirilməlidir", True),
    (20, "ГОСТ Р 55724-2013", "QNM. Qaynaq birləşmələri. Ultrasəs üsulları",
         "UT — ultrasəs müayinə", False),
    (21, "ГОСТ Р 56512-2015", "QNM. Maqnit-toz üsulu. Tipik texnoloji proseslər",
         "MT — maqnit-toz müayinə", False),
]


def _check_fonts() -> None:
    missing = [str(p) for p in (FONT_REGULAR, FONT_BOLD) if not p.exists()]
    if missing:
        sys.exit(
            "ERROR: required font file(s) missing:\n  "
            + "\n  ".join(missing)
            + "\nPlace the DejaVu Sans TTFs in assets/fonts/ (see assets/fonts/README.md)."
        )


# --- DOCX ------------------------------------------------------------------------------

def build_docx() -> None:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor

    FONT_NAME = "DejaVu Sans"

    def _shade(cell, hex_fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def _set_run_font(run, size: int, bold: bool = False,
                      color: RGBColor | None = None) -> None:
        run.font.name = FONT_NAME
        # Ensure the font is applied to complex/east-asian script runs too.
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), FONT_NAME)
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color

    doc = Document()

    # Default document font.
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)

    # Title.
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.add_run(TITLE), size=14, bold=True)
    doc.add_paragraph()  # spacer

    # Intro.
    intro_p = doc.add_paragraph()
    intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_run_font(intro_p.add_run(INTRO), size=11)
    doc.add_paragraph()  # spacer

    # Table.
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for cell, text in zip(hdr, COL_HEADERS):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade(cell, "2F5496")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(text), size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for no, code, full_name, method, _uncertain in STANDARDS:
        cells = table.add_row().cells
        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p0.add_run(str(no)), size=11)

        p1 = cells[1].paragraphs[0]
        _set_run_font(p1.add_run(code), size=11, bold=True)
        _set_run_font(p1.add_run(f" — {full_name}"), size=11)

        p2 = cells[2].paragraphs[0]
        _set_run_font(p2.add_run(method), size=11)

    # Approximate column widths (python-docx sets per-cell width).
    from docx.shared import Cm
    widths = (Cm(1.0), Cm(9.5), Cm(6.5))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w

    doc.add_paragraph()
    doc.add_paragraph()

    # Signature block (left-aligned placeholder lines).
    for label in SIGNATURE_LABELS:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_run_font(sp.add_run(f"{label}: ______________________________"), size=11)

    doc.save(str(OUT_DOCX))


# --- PDF -------------------------------------------------------------------------------

def build_pdf() -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_REGULAR)))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_BOLD)))
    pdfmetrics.registerFontFamily(
        "DejaVu", normal="DejaVu", bold="DejaVu-Bold",
        italic="DejaVu", boldItalic="DejaVu-Bold",
    )

    title_style = ParagraphStyle(
        "TitleAZ", fontName="DejaVu-Bold", fontSize=14, leading=18,
        alignment=TA_CENTER, spaceAfter=14,
    )
    intro_style = ParagraphStyle(
        "IntroAZ", fontName="DejaVu", fontSize=10.5, leading=15,
        alignment=TA_JUSTIFY, spaceAfter=12,
    )
    cell_style = ParagraphStyle(
        "CellAZ", fontName="DejaVu", fontSize=9.5, leading=12.5, alignment=TA_LEFT,
    )
    cell_center = ParagraphStyle(
        "CellCenterAZ", fontName="DejaVu", fontSize=9.5, leading=12.5, alignment=TA_CENTER,
    )
    hdr_style = ParagraphStyle(
        "HdrAZ", fontName="DejaVu-Bold", fontSize=10, leading=13,
        alignment=TA_CENTER, textColor=colors.white,
    )
    sig_style = ParagraphStyle(
        "SigAZ", fontName="DejaVu", fontSize=10.5, leading=22, alignment=TA_LEFT,
    )

    doc = SimpleDocTemplate(
        str(OUT_PDF), pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title="Add-portal standartlar sorğusu",
    )

    story: list = [
        Paragraph(TITLE, title_style),
        Paragraph(INTRO, intro_style),
    ]

    # Table data: header + 21 rows, cells wrapped as Paragraphs so long AZ titles flow.
    data = [[Paragraph(h, hdr_style) for h in COL_HEADERS]]
    for no, code, full_name, method, _uncertain in STANDARDS:
        name_html = f"<b>{code}</b> — {full_name}"
        data.append([
            Paragraph(str(no), cell_center),
            Paragraph(name_html, cell_style),
            Paragraph(method, cell_style),
        ])

    # A4 content width ~ 174 mm; split narrow / wide / medium.
    col_widths = [12 * mm, 100 * mm, 62 * mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5496")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8EAADB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF1F9")]),
    ]))
    story.append(table)

    story.append(Spacer(1, 18 * mm))
    for label in SIGNATURE_LABELS:
        story.append(Paragraph(f"{label}: ______________________________", sig_style))

    doc.build(story)


# --- Verify ----------------------------------------------------------------------------

def verify() -> None:
    from docx import Document

    ok = True

    # DOCX checks.
    if not OUT_DOCX.exists() or OUT_DOCX.stat().st_size == 0:
        print(f"  [FAIL] {OUT_DOCX.name} missing or empty")
        ok = False
    else:
        d = Document(str(OUT_DOCX))
        rows = len(d.tables[0].rows) if d.tables else 0
        has_title = any(TITLE in p.text for p in d.paragraphs)
        print(f"  [ok] {OUT_DOCX.name}  {OUT_DOCX.stat().st_size:,} bytes, "
              f"table rows={rows}, title={'yes' if has_title else 'NO'}")
        if rows != len(STANDARDS) + 1 or not has_title:
            ok = False

    # PDF checks.
    if not OUT_PDF.exists() or OUT_PDF.stat().st_size < 3000:
        print(f"  [FAIL] {OUT_PDF.name} missing or too small")
        ok = False
    else:
        head = OUT_PDF.read_bytes()[:5]
        good = head.startswith(b"%PDF")
        print(f"  [{'ok' if good else 'FAIL'}] {OUT_PDF.name}  "
              f"{OUT_PDF.stat().st_size:,} bytes, header={head!r}")
        if not good:
            ok = False

    flagged = [(c, n) for _no, c, n, _m, unc in STANDARDS if unc]
    if flagged:
        print("\n  Operator review — GOST titles not fully confirmed:")
        for code, name in flagged:
            print(f"    * {code}  ({name})")

    if not ok:
        sys.exit("\nVerification FAILED — see messages above.")
    print("\nDone. Both documents generated and verified.")


def main() -> None:
    _check_fonts()
    print("Generating official add-portal request letter...")
    build_docx()
    build_pdf()
    verify()


if __name__ == "__main__":
    main()
