"""The native ``.docx`` path must be byte-identical after ADR-003.

ADR-003 changes how PDFs are read. It must change **nothing** about a report that
arrives as a real ``.docx``, and "nothing" is not a claim to make by reading the
diff: ``_read_one`` gained a suffix branch, ``parse_cache`` gained a schema bump
and a store-skip rule, and any of those could have leaked onto the path that
carries the overwhelming majority of the corpus (~3,900 native documents against
660 PDF-derived ones).

So the whole native path is exercised end to end — build ``.docx`` files, read them
through ``docx_reader.read_many``, aggregate, and export the JSON sidecar — and the
result is compared against ``tests/data/native_docx_sidecar.json``.

**How that file was produced:** the same generator was run against a git worktree
holding the pre-ADR-003 code, with every file this work did *not* touch copied in
from the working tree. So the golden is exactly "today's pipeline minus this
change", and a mismatch means the change reached somewhere it had no business
reaching. It was byte-identical when the golden was captured (2026-09-07).

If this test fails after an intentional change to the sidecar schema, regenerate
the golden **in a separate commit** so the diff shows a human deciding to change
the sidecar, rather than a reader refactor changing it by accident.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from monthly_report import (config, docx_reader, document_aggregator,
                            json_exporter, translator as translator_mod)

_ROOT = Path(__file__).resolve().parent.parent
_GOLDEN = json.loads(
    (Path(__file__).parent / "data" / "native_docx_sidecar.json").read_text(encoding="utf-8"))


def _make_report(path: Path, report_no: str, welds) -> Path:
    """A minimal but realistic native report: header block, weld table, signature."""
    doc = Document()
    header = doc.add_table(rows=8, cols=2)
    for i, (label, value) in enumerate([
            ("Report No:", report_no), ("Date:", "10.09.2025"),
            ("Client:", "KHAZAR FABRICATION"), ("Acc Standard:", "ASME B31.3"),
            ("Item Inspection:", "Pipe"), ("Material:", "CS X60"),
            ("Procedure:", "P-01"), ("Drw:", "D-01")]):
        header.cell(i, 0).text = label
        header.cell(i, 1).text = value

    table = doc.add_table(rows=1 + len(welds), cols=4)
    for j, head in enumerate(["Weld No", "Diam./Thick.", "Item", "Result"]):
        table.cell(0, j).text = head
    for i, row in enumerate(welds, 1):
        for j, value in enumerate(row):
            table.cell(i, j).text = value

    tail = doc.add_table(rows=2, cols=2)
    tail.cell(0, 0).text = "Equipment Used:"
    tail.cell(0, 1).text = "EPOCH LT"
    tail.cell(1, 0).text = "Name:"
    tail.cell(1, 1).text = "R. Aliyev"
    doc.save(str(path))
    return path


@pytest.fixture(scope="module")
def sidecar(tmp_path_factory) -> dict:
    """Run the corpus through the real native path and return the sidecar payload."""
    settings_file = _ROOT / "config" / "settings.yaml"
    if not settings_file.is_file():
        pytest.skip("config/settings.yaml not present (it is gitignored)")

    work = tmp_path_factory.mktemp("native")
    source = work / "input" / "MPI"
    source.mkdir(parents=True)
    # Six reports: the client volume gate drops a batch smaller than five, and two
    # of them reject so both verdict branches are covered.
    for i in range(1, 7):
        report_no = f"KF-MT-26-{i:03d}"
        welds = ([("W1", "6", "Pipe", "Accept"), ("W2", "8", "Pipe", "Accept")]
                 if i % 3 else [("W1", "6", "Pipe", "Reject")])
        _make_report(source / f"{report_no} (2165) (ASME B31.3)_en.docx",
                     report_no, welds)

    settings = config.load_settings(str(settings_file))
    settings.input_dir = str(work / "input")
    settings.output_dir = str(work / "output")
    settings.target_month = "2025-09"
    # The golden carries a fake VÖEN, not the live one: no real tax ID in the repo.
    settings.raw["clients"]["KHF"]["voen"] = "1000000002"

    paths = sorted((work / "input").rglob("*.docx"))
    cache, errors, textless = docx_reader.read_many(paths, max_workers=1)
    assert not errors and not textless

    trans = translator_mod.Translator(
        translator_mod.load_dictionary(str(_ROOT / "translations" / "az.json")))
    report = document_aggregator.build_monthly_report(
        settings, trans, method="MT", client_code="KHF", seq=1, cache=cache)
    assert report.skipped == [], report.skipped
    return json_exporter.build_report_json(report, settings, trans)


def test_the_sidecar_is_unchanged(sidecar) -> None:
    assert sidecar == _GOLDEN


def test_every_row_survived(sidecar) -> None:
    """Guards the guard: an empty batch would also equal an empty golden."""
    assert sidecar["total_records"] == 6
    assert len(sidecar["items_table"]) == 6


def test_the_verdicts_are_the_ones_the_documents_carry(sidecar) -> None:
    """The field the whole of ADR-003 is about, on the path it must not touch."""
    assert [row["result"] for row in sidecar["items_table"]] == [
        "Yararlı", "Yararlı", "Yarasız", "Yararlı", "Yararlı", "Yarasız"]


def test_no_disagreement_is_produced_for_a_docx(sidecar) -> None:
    """The cross-check is a PDF-only mechanism; a .docx batch must not enter it."""
    assert docx_reader.read_many.last_disagreements == []
