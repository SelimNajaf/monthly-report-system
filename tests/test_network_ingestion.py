"""STAGE 0-LAN must stage new target-year reports off the share — read-only, deduped,
and never crashing when the share is offline.

``services.network_ingestion.sync_from_lan`` copies fresh ``.docx``/``.pdf`` reports
from the configured LAN roots into ``input/UnSorted/`` before the local pipeline runs.
The guarantees pinned here:

* **Year scoping** — a wrong-year folder (``UT-2025`` for a 2026 target) is pruned;
  only the target year's files are staged.
* **Dedup** — a filename already under ``input/`` is skipped with NO transfer; a
  differently-named byte-identical file is copied then removed via the content-hash
  net (both a ``.parse_cache`` digest and a within-run twin).
* **Read-only** — every source file survives with an unchanged mtime; nothing on the
  "share" is moved or deleted.
* **Graceful offline** — every root unreachable ⇒ ``unreachable_all`` with no raise.

A local ``tmp_path`` tree stands in for the UNC share (the walk is identical on a
local path), so no live network is ever touched.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

import discover_network_share as dns
from monthly_report import config
from monthly_report.services import network_ingestion as ni

#: The engagement cutoff exercised by the internal-date engine tests.
CUTOFF = dt.date(2025, 12, 1)

_SETTINGS_YAML = Path(__file__).resolve().parents[1] / "config" / "settings.yaml"


# --- fixtures / builders ----------------------------------------------------------------


def _real_registries():
    """The LIVE ``clients`` + ``ndt_methods`` blocks out of ``config/settings.yaml``.

    The filename classifier is config-driven, so the fake settings the sync tests use
    must carry the REAL registries — a hand-written copy would drift the day someone
    adds a client and the gate would silently stop matching production.
    """
    raw = config.load_settings(_SETTINGS_YAML).raw
    return raw.get("clients", {}), raw.get("ndt_methods", {})


@pytest.fixture(scope="module")
def real_settings():
    """The full production settings object (classifier tests read its registries)."""
    return config.load_settings(_SETTINGS_YAML)


def _settings(tmp_path, sources, *, target_month="2026-08"):
    """A minimal stand-in for config.Settings carrying only what the sync reads."""
    input_dir = tmp_path / "input"
    input_dir.mkdir(exist_ok=True)
    clients, ndt_methods = _real_registries()
    return SimpleNamespace(
        input_dir=input_dir,
        project_root=tmp_path,
        target_month=target_month,
        raw={
            "network_ingestion": {"lan_sources": [str(s) for s in sources]},
            "performance": {"parse_cache_dir": ".parse_cache"},
            "clients": clients,
            "ndt_methods": ndt_methods,
        },
    )


def _write(path, data=b"report-bytes"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def _unsorted_names(settings):
    d = settings.input_dir / "UnSorted"
    return sorted(p.name for p in d.iterdir()) if d.is_dir() else []


def _make_dated_docx(path, date_str):
    """Write a real .docx whose ``Date:`` cell carries ``date_str`` (``""`` = no date).

    A 1-row, 2-col table — label cell + value cell — is exactly what
    ``extractor.extract_meta`` reads to recover the internal inspection date.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Date:"
    table.cell(0, 1).text = date_str
    doc.save(str(path))
    return path


def _manifest_path(tmp_path):
    return tmp_path / ".state" / "network_sync_manifest.json"


# --- pure helpers -----------------------------------------------------------------------


class TestPureHelpers:
    @pytest.mark.parametrize("month, year", [
        ("2026-08", "2026"), ("2026-1", "2026"), ("1999-12", "1999"),
        ("", ""), ("garbage", ""),
    ])
    def test_target_year(self, month, year):
        assert ni._target_year(month) == year

    def test_share_label_unc(self):
        got = ni._share_label(r"\\lan-host\Shared Docs\UT")
        assert got == r"\\lan-host\Shared Docs"

    def test_share_label_local_passthrough(self, tmp_path):
        # A non-UNC path is returned unchanged (test/local share).
        assert ni._share_label(str(tmp_path)) == str(tmp_path)

    def test_host_of(self):
        assert ni._host_of(r"\\lan-host\Share\UT") == "lan-host"
        assert ni._host_of("C:\\local\\path") is None


# --- year scoping -----------------------------------------------------------------------


class TestYearScoping:
    def test_only_target_year_is_staged(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "UBOC" / "U-UT-26-001.docx")
        _write(share / "UT-2025" / "UBOC" / "U-UT-25-001.docx")   # wrong year — pruned
        settings = _settings(tmp_path, [share])

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.staged == ["U-UT-26-001.docx"]
        assert _unsorted_names(settings) == ["U-UT-26-001.docx"]

    def test_non_year_folders_descended(self, tmp_path):
        # A file under a non-year-named subfolder is never pruned by the year gate.
        share = tmp_path / "share" / "MPI"
        _write(share / "2026" / "AZFEN" / "Report" / "AZF-MT-26-009.docx")
        settings = _settings(tmp_path, [share])

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.staged == ["AZF-MT-26-009.docx"]

    def test_default_month_uses_settings_target(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "X" / "U-UT-26-011.docx")
        _write(share / "UT-2024" / "X" / "U-UT-24-011.docx")
        settings = _settings(tmp_path, [share], target_month="2026-01")

        # Caller passes settings.target_month when no --month is given.
        result = ni.sync_from_lan(settings, settings.target_month, timeout=None)

        assert result.staged == ["U-UT-26-011.docx"]


# --- deduplication ----------------------------------------------------------------------


class TestDedup:
    def test_existing_filename_skipped_without_transfer(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "U-UT-26-050.docx", b"share-copy")
        settings = _settings(tmp_path, [share])
        # A same-stem file already settled under input/ (different bytes).
        _write(settings.input_dir / "UT" / "U-UT-26-050.docx", b"local-copy")

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.staged == []
        assert result.skipped_dupes == 1
        assert _unsorted_names(settings) == []

    def test_content_hash_from_parse_cache_removes_copy(self, tmp_path):
        data = b"identical-report-payload"
        digest = hashlib.sha256(data).hexdigest()
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "KF-UT-26-088.docx", data)
        settings = _settings(tmp_path, [share])
        # A .parse_cache entry proves this content is already known locally.
        cache = tmp_path / ".parse_cache"
        cache.mkdir()
        (cache / f"{digest}.json").write_text("{}", encoding="utf-8")

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.staged == []
        assert result.skipped_dupes == 1
        assert _unsorted_names(settings) == []          # copied then removed

    def test_within_run_identical_twin_removed(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "U-UT-26-061.docx", b"same")
        _write(share / "UT-2026" / "C" / "U-UT-26-062.docx", b"same")
        settings = _settings(tmp_path, [share])

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        # One copy survives; its identical twin is caught by the content net.
        assert len(result.staged) == 1
        assert result.skipped_dupes == 1
        assert len(_unsorted_names(settings)) == 1

    def test_second_run_is_idempotent(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "U-UT-26-070.docx", b"payload")
        settings = _settings(tmp_path, [share])

        first = ni.sync_from_lan(settings, "2026-08", timeout=None)
        second = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert first.staged_count == 1
        assert second.staged_count == 0                 # already staged -> no re-copy
        # The second run answers from the ingestion registry's (size, mtime)
        # fingerprint, so it never reaches the filename-index path that
        # `skipped_dupes` counts. Same outcome, one dict lookup instead of a walk.
        assert second.unchanged == 1
        assert second.skipped_dupes == 0
        assert second.checked == 1

    def test_filename_skip_still_counts_without_a_registry(self, tmp_path):
        """A cold registry (first run over an already-populated input/) still dedups.

        The O(1) path is an accelerator, not the correctness boundary — with no
        fingerprint on file, the original filename index is what stops the transfer.
        """
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "U-UT-26-071.docx", b"payload")
        settings = _settings(tmp_path, [share])
        _write(settings.input_dir / "UT" / "U-UT-26-071.docx", b"payload")

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.staged == []
        assert result.skipped_dupes == 1
        assert result.unchanged == 0
        # ...and it is fingerprinted on the way out, so the NEXT run is O(1).
        assert ni.sync_from_lan(settings, "2026-08", timeout=None).unchanged == 1

    # --- A4-V2: the MODIFIED branch dedups too --------------------------------------

    def test_modified_identical_to_a_live_twin_is_dropped(self, tmp_path):
        """Two live copies + unchanged bytes ⇒ no `-lan` artifact.

        The share file is MODIFIED (fingerprint moved), but its new bytes already
        exist locally. With two live copies there is no single one to replace, so the
        old code staged a `-lan` twin blind; now the content net catches it.
        """
        share = tmp_path / "share" / "UT"
        src = _write(share / "UT-2026" / "C" / "U-UT-26-080.docx", b"first-payload")
        settings = _settings(tmp_path, [share])
        ni.sync_from_lan(settings, "2026-08", timeout=None)        # stages UnSorted/
        # A second LIVE copy carrying the bytes the share is about to hold.
        _write(settings.input_dir / "UT" / "U-UT-26-080.docx", b"second-payload-v2")
        src.write_bytes(b"second-payload-v2")                      # share now MODIFIED

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.skipped_dupes == 1
        assert result.modified == []
        assert _unsorted_names(settings) == ["U-UT-26-080.docx"]   # no -lan twin

    def test_modified_with_new_bytes_still_stages_the_lan_twin(self, tmp_path):
        """Genuinely different content keeps the non-clobbering twin — no data loss."""
        share = tmp_path / "share" / "UT"
        src = _write(share / "UT-2026" / "C" / "U-UT-26-081.docx", b"first-payload")
        settings = _settings(tmp_path, [share])
        ni.sync_from_lan(settings, "2026-08", timeout=None)
        _write(settings.input_dir / "UT" / "U-UT-26-081.docx", b"an-unrelated-copy")
        src.write_bytes(b"genuinely-new-third-payload")

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.skipped_dupes == 0
        assert result.modified == ["U-UT-26-081-lan.docx"]
        assert _unsorted_names(settings) == ["U-UT-26-081-lan.docx",
                                             "U-UT-26-081.docx"]

    def test_modified_with_one_live_copy_is_replaced_in_place(self, tmp_path):
        """The `replacing` path is untouched by the dedup: overwrite, never drop."""
        share = tmp_path / "share" / "UT"
        src = _write(share / "UT-2026" / "C" / "U-UT-26-082.docx", b"first-payload")
        settings = _settings(tmp_path, [share])
        ni.sync_from_lan(settings, "2026-08", timeout=None)
        src.write_bytes(b"refreshed-payload-v2")

        result = ni.sync_from_lan(settings, "2026-08", timeout=None)

        assert result.skipped_dupes == 0
        assert result.modified == ["U-UT-26-082.docx"]
        assert _unsorted_names(settings) == ["U-UT-26-082.docx"]
        staged = settings.input_dir / "UnSorted" / "U-UT-26-082.docx"
        assert staged.read_bytes() == b"refreshed-payload-v2"

    def test_modified_identical_to_a_settled_copy_names_the_folder(self, tmp_path):
        """A SKIPPED_REPORTS sibling counts as a match, and the log says where.

        The operator needs to see WHY a copy was dropped when the twin sits in a
        parked tree — a quarantined report must not silently return to UnSorted/.
        """
        share = tmp_path / "share" / "UT"
        src = _write(share / "UT-2026" / "C" / "U-UT-26-083.docx", b"first-payload")
        settings = _settings(tmp_path, [share])
        ni.sync_from_lan(settings, "2026-08", timeout=None)
        # The staged copy is quarantined by the operator and updated on the share.
        (settings.input_dir / "UnSorted" / "U-UT-26-083.docx").unlink()
        _write(settings.input_dir / "SKIPPED_REPORTS" / "U-UT-26-083.docx",
               b"quarantined-payload-v2")
        src.write_bytes(b"quarantined-payload-v2")

        logs = []
        result = ni.sync_from_lan(settings, "2026-08", timeout=None, log=logs.append)

        assert result.skipped_dupes == 1
        assert result.modified == []
        assert _unsorted_names(settings) == []                     # nothing came back
        assert any("Skipped duplicate content U-UT-26-083.docx "
                   "(matches SKIPPED_REPORTS/U-UT-26-083.docx)" in line
                   for line in logs), logs


# --- read-only guarantee ----------------------------------------------------------------


class TestReadOnly:
    def test_share_files_untouched(self, tmp_path):
        share = tmp_path / "share" / "UT"
        src_a = _write(share / "UT-2026" / "C" / "U-UT-26-021.docx", b"aaa")
        src_b = _write(share / "UT-2026" / "C" / "KF-UT-26-022.pdf", b"%PDF-bbb")
        before = {p: (p.read_bytes(), p.stat().st_mtime) for p in (src_a, src_b)}
        settings = _settings(tmp_path, [share])

        ni.sync_from_lan(settings, "2026-08", timeout=None)

        for p, (data, mtime) in before.items():
            assert p.exists(), "source file must not be moved/deleted"
            assert p.read_bytes() == data
            assert p.stat().st_mtime == pytest.approx(mtime)


# --- graceful offline -------------------------------------------------------------------


class TestOffline:
    def test_all_roots_unreachable(self, monkeypatch, tmp_path):
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "C" / "U-UT-26-031.docx")
        settings = _settings(tmp_path, [share])
        # Probe never answers reachable -> the root is treated as offline.
        monkeypatch.setattr(ni, "_reachable_within", lambda root, timeout: None)

        result = ni.sync_from_lan(settings, "2026-08")   # default timeout -> probe path

        assert result.unreachable_all is True
        assert result.staged == []
        assert result.roots_offline == 1

    def test_dead_host_probed_once(self, monkeypatch, tmp_path):
        # Two roots on one dead host: only the FIRST is probed, the rest inherit
        # "offline" without another timeout.
        calls = []

        def fake_probe(root, timeout):
            calls.append(str(root))
            return None

        monkeypatch.setattr(ni, "_reachable_within", fake_probe)
        settings = _settings(tmp_path, [r"\\lan-host\Share\UT",
                                        r"\\lan-host\Share\RT"])
        result = ni.sync_from_lan(settings, "2026-08")

        assert len(calls) == 1                          # second root not re-probed
        assert result.roots_offline == 2
        assert result.unreachable_all is True

    def test_no_sources_is_noop(self, tmp_path):
        settings = _settings(tmp_path, [])
        result = ni.sync_from_lan(settings, "2026-08")
        assert result.roots_scanned == 0
        assert result.unreachable_all is False          # not configured != outage
        assert result.staged == []


# --- CLI month normalisation (main.py) --------------------------------------------------


class TestShortMonthForm:
    def test_yy_mm_form(self):
        import main
        assert main._normalize_month("26-08") == "2026-08"
        assert main._normalize_month("26-6") == "2026-06"

    def test_full_forms_still_work(self):
        import main
        assert main._normalize_month("2026-08") == "2026-08"
        assert main._normalize_month("JUN_2026") == "2026-06"


# --- cutoff engine (sync_since) ---------------------------------------------------------


class TestCutoffEngine:
    def test_legacy_before_cutoff_skipped(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2025" / "C" / "U-UT-25-900.docx", "10.11.2025")  # legacy
        _make_dated_docx(share / "UT-2025" / "C" / "U-UT-25-950.docx", "05.12.2025")  # keep
        _make_dated_docx(share / "UT-2026" / "C" / "U-UT-26-001.docx", "14.01.2026")  # keep
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.legacy_skipped == 1
        assert set(r.staged) == {"U-UT-25-950.docx", "U-UT-26-001.docx"}
        assert set(_unsorted_names(settings)) == {"U-UT-25-950.docx", "U-UT-26-001.docx"}

    def test_undetermined_date_is_kept(self, tmp_path):
        # A .docx with no parseable Date: value is staged, never dropped.
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2025" / "C" / "U-UT-25-777.docx", "")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["U-UT-25-777.docx"]
        assert r.legacy_skipped == 0

    def test_docx_twin_preferred_over_pdf(self, tmp_path):
        share = tmp_path / "share" / "DP"
        _make_dated_docx(share / "DP-2026" / "C" / "U-DP-26-002.docx", "15.12.2025")
        _write(share / "DP-2026" / "C" / "U-DP-26-002.pdf", b"%PDF-twin")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["U-DP-26-002.docx"]
        assert r.twin_skipped == 1
        assert _unsorted_names(settings) == ["U-DP-26-002.docx"]

    def test_lone_pdf_is_copied(self, tmp_path):
        # A .pdf with no .docx twin can't be date-gated; it is staged as-is.
        share = tmp_path / "share" / "RT"
        _write(share / "RT-2026" / "C" / "U-RT-26-500.pdf", b"%PDF-lone")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["U-RT-26-500.pdf"]
        assert r.twin_skipped == 0

    def test_manifest_dedup_idempotent(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2026" / "C" / "U-UT-26-070.docx", "20.12.2025")
        settings = _settings(tmp_path, [share])
        mpath = _manifest_path(tmp_path)

        first = ni.sync_since(settings, CUTOFF, manifest_path=mpath, timeout=None)
        second = ni.sync_since(settings, CUTOFF, manifest_path=mpath, timeout=None)

        assert first.staged_count == 1
        assert second.staged_count == 0
        assert second.up_to_date >= 1
        data = json.loads(mpath.read_text(encoding="utf-8"))
        assert data["cutoff"] == CUTOFF.isoformat()
        assert len(data["entries"]) == 1                 # the staged report's hash

    def test_min_year_folder_pruned(self, tmp_path):
        # A pre-cutoff-year folder is pruned before any .docx is opened.
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2024" / "C" / "U-UT-24-800.docx", "10.06.2024")
        _make_dated_docx(share / "UT-2026" / "C" / "U-UT-26-800.docx", "10.06.2026")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["U-UT-26-800.docx"]
        assert r.legacy_skipped == 0                     # 2024 folder never descended


# --- filename classifier ----------------------------------------------------------------
#
# Every sample below is a real path taken verbatim from file_paths_ex.md — the
# operator's inventory of the live share. They are the contract: each of these MUST
# be recognised as a report, across all seven NDT methods and both name orders.

#: Every sample filename in file_paths_ex.md, grouped by its share method root.
SHARE_SAMPLES = [
    # Dye Penetrant — client-first, weld token, parenthesised job no, twins
    "AZF-DP-W1-25-39.docx",
    "AZF-DP-W1-25-39.pdf",
    "G-DP-25-001.docx",
    "G-DP-25-001.pdf",
    "G-DP-25-009R1.docx",
    "KF-DP-25-001(1896).docx",
    "KF-DP-25-001(1896).pdf",
    "NOV-DP-25-001.pdf",
    "NOV-DP-25-001.docx",
    "U-DP-25-001.pdf",
    "U-DP-25-001.docx",
    # Dye Penetrant — SDLN method-first with the job/size/date tail
    "DP-SDLN-25-004(F1798)6”-15E440-03-A_12.09.2025.pdf",
    "DP-SDLN-25-005(F1798)6”-15E440-02-A_16.09.2025.docx",
    # Hardness Test (HrT) — kept; Heat Treatment is dropped at the source level
    "KF-HrT-25-001 (1901).pdf",
    "KF-HrT-25-001 (1901).docx",
    "KF-HrT-26-001 (2162).docx",
    # MPI / MT
    "G-MT-25-001.docx",
    "G-MT-25-001.pdf",
    "MPI-SDLN-25-112(F1809)6”x7.11_6”-071001-01-A_30.09.2025.docx",
    "MPI-SDLN-25-111(F1785)6”x7.11_6”-02D122-03-B_30.09.2025.docx",
    "U-MT-25-002 (Dashgil 2).docx",
    "U-MT-25-003 (Platform).docx",
    # PMI — both the operator prefix (KF) and the registry code (UBOC) forms
    "KF-PMI-25-03.docx",
    "KF-PMI-25-02.docx",
    "UBOC-PMI-25-001.docx",
    "UBOC-PMI-25-001.pdf",
    # UT — including the legacy .doc exports
    "NOV-UT-25-001.pdf",
    "NOV-UT-25-001.docx",
    "UT-SDLN-25-044ASR(F1827)30_x 38.0_W-2ASR_08.10.2025.doc",
    "UT-SDLN-25-045(F1827)30_x 38.0_W-3_08.10.2025.doc",
    # X-ray / RT
    "KF-RT-25-004 (KF 1866).docx",
    "KF-RT-25-003 ASR(KF 1866).docx",
    "U-RT-25-126.docx",
    "U-RT-25-125 (Dashgil 2).docx",
    # Ferrite content / FC
    "KF-FC-25-02.pdf",
    "KF-FC-25-02.docx",
]

#: Names that must NEVER be staged.
JUNK_SAMPLES = [
    "~$KF-DP-25-001.docx",              # Word lock file
    "~$G-MT-25-001.docx",
    ".~lock.U-UT-25-001.docx#",         # editor lock dropping
    "Report Forms.docx",                # administrative template
    "Sertifikat.pdf",                   # certificate, not a report
    "Price list 2026.docx",
    "AZFEN 2026 monthly summary.docx",  # summary doc, no report serial
    "notes.txt",                        # not a report extension
    "MT-26-001.docx",                   # method-first with NO client token
    "XYZ-DP-26-001.docx",               # unregistered client prefix
    "IMG_20260114.jpg",
    "Thumbs.db",
]


class TestReportFilenameClassifier:
    @pytest.mark.parametrize("name", SHARE_SAMPLES)
    def test_every_share_sample_is_a_report(self, name, real_settings):
        assert dns.is_valid_report_filename(name, real_settings), name

    @pytest.mark.parametrize("name", JUNK_SAMPLES)
    def test_junk_is_rejected(self, name, real_settings):
        assert not dns.is_valid_report_filename(name, real_settings), name

    def test_full_unc_path_basename_is_classified(self, real_settings):
        # The walk hands the classifier a bare entry.name; taking it off a UNC path
        # must give the same verdict.
        path = ("\\\\lan-host\\Shared Docs\\Dye Penetrant\\DP-2025\\GEMZA\\"
                "G-DP-25-001.docx")
        assert dns.is_valid_report_filename(Path(path).name, real_settings)

    def test_retired_client_still_recognised(self, real_settings):
        # SIENMAR (S) is active:false — it is RECOGNISED here and quarantined by the
        # pipeline downstream; dropping it at the LAN gate would lose that audit trail.
        assert dns.is_valid_report_filename("S-UT-26-001.docx", real_settings)

    def test_empty_and_none_are_safe(self, real_settings):
        assert not dns.is_valid_report_filename("", real_settings)
        assert not dns.is_valid_report_filename(None, real_settings)

    def test_every_method_root_is_covered(self, real_settings):
        # Guards the sample list itself: all seven live methods must be exercised.
        from monthly_report import method_registry as mr
        from monthly_report.report_identity import report_key_of
        methods = {mr.detect_method(report_key_of(n), real_settings)
                   for n in SHARE_SAMPLES}
        assert {"DP", "MT", "MPI", "RT", "UT", "PMI", "FC", "HRT"} & methods == methods
        assert len(methods) >= 7


class TestShareSpellingVariants:
    """Alternate spellings found on the live share during the cutover audit.

    Both were silently rejected before they were registered: 57 SDLN reports dated
    December 2025 and 84 AZFEN Ferrite reports. Pinned here so a registry edit that
    drops either token fails loudly instead of quietly shrinking the corpus.
    """

    @pytest.mark.parametrize("name", [
        "RT-SDL-25-333(F1840)6_x7.11_6”-03D121-61-A_W-1_03.12.2025.docx",
        "RT-SDL-25-001(F1752)30'x38_W-83_01.09.2025.pdf",
    ])
    def test_short_sdl_client_token_recognised(self, name, real_settings):
        # The share writes both "SDLN" and "SDL" inside the SAME RT-2025/SDLN folder.
        from monthly_report import method_registry as mr
        from monthly_report.report_identity import report_key_of
        assert dns.is_valid_report_filename(name, real_settings), name
        assert mr.resolve_client(report_key_of(name), real_settings)[0] == "SDLN"

    def test_fcm_resolves_to_the_ferrite_method(self, real_settings):
        from monthly_report import method_registry as mr
        from monthly_report.report_identity import report_key_of
        key = report_key_of("AZF-FCM-W1-25-01.docx")
        assert dns.is_valid_report_filename("AZF-FCM-W1-25-01.docx", real_settings)
        assert mr.detect_method(key, real_settings) == "FCM"
        assert mr.method_folder("FCM", real_settings) == "FERRITE"
        assert mr.detect_client(key, real_settings) == "AZF"
        # FCM must not be swallowed by the shorter FC token, and inherits its default.
        assert mr.method_default_standard(real_settings, "FCM") == "AWS A4.2M"

    def test_fc_still_resolves_after_fcm_was_added(self, real_settings):
        from monthly_report import method_registry as mr
        assert mr.detect_method("KF-FC-25-02", real_settings) == "FC"
        assert mr.method_folder("FC", real_settings) == "FERRITE"


class TestUtilityFolders:
    @pytest.mark.parametrize("name", [
        "Templates", "TEMPLATES", "Old reports", "Archive", "Backup 2024",
        "Report Forms", "Certificates", "Equipment", "Procedure", "Draft",
    ])
    def test_utility_folders_detected(self, name):
        assert dns.is_utility_folder(name)

    @pytest.mark.parametrize("name", [
        "UBOC", "KHAZAR", "GEMZA", "SDLN", "AZFEN", "NOV", "DP-2026", "2026",
    ])
    def test_client_and_year_folders_are_not_utility(self, name):
        assert not dns.is_utility_folder(name)


class TestSerialAndFilenameDate:
    @pytest.mark.parametrize("name, expected", [
        ("AZF-DP-W1-25-39.docx", (25, 39)),
        ("G-DP-25-009R1.docx", (25, 9)),
        ("KF-DP-25-001(1896).pdf", (25, 1)),
        ("DP-SDLN-25-005(F1798)6”-15E440-02-A_16.09.2025.docx", (25, 5)),
        ("KF-HrT-26-001 (2162).docx", (26, 1)),
        ("UT-SDLN-25-044ASR(F1827)30_x 38.0_W-2ASR_08.10.2025.doc", (25, 44)),
        ("KF-FC-25-02.docx", (25, 2)),
        ("Report Forms.docx", None),
        ("Sertifikat.pdf", None),
    ])
    def test_report_serial(self, name, expected):
        assert dns.report_serial(name) == expected

    @pytest.mark.parametrize("name, expected", [
        ("DP-SDLN-25-004(F1798)6”-15E440-03-A_12.09.2025.pdf",
         dt.date(2025, 9, 12)),
        ("MPI-SDLN-25-112(F1809)6”x7.11_6”-071001-01-A_30.09.2025.docx",
         dt.date(2025, 9, 30)),
        ("UT-SDLN-25-045(F1827)30_x 38.0_W-3_08.10.2025.doc", dt.date(2025, 10, 8)),
        ("G-DP-25-001.docx", None),                 # no date in the name
        ("KF-DP-25-001(1896).docx", None),          # a job number is not a date
    ])
    def test_filename_date(self, name, expected):
        assert dns.filename_date(name) == expected

    def test_impossible_date_is_ignored(self):
        assert dns.filename_date("U-UT-25-001_32.13.2025.pdf") is None


# --- ingestion filter, end to end --------------------------------------------------------


class TestIngestionFilter:
    def test_junk_and_lock_files_never_staged(self, tmp_path):
        share = tmp_path / "share" / "DP"
        _make_dated_docx(share / "DP-2026" / "KHAZAR" / "KF-DP-26-001.docx",
                         "10.01.2026")
        _write(share / "DP-2026" / "KHAZAR" / "~$KF-DP-26-001.docx", b"lock")
        _write(share / "DP-2026" / "KHAZAR" / "Report Forms.docx", b"admin")
        _write(share / "DP-2026" / "KHAZAR" / "Sertifikat.pdf", b"%PDF-cert")
        _write(share / "DP-2026" / "KHAZAR" / "notes.txt", b"text")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["KF-DP-26-001.docx"]
        assert _unsorted_names(settings) == ["KF-DP-26-001.docx"]
        assert r.rejected == 4          # lock file, admin doc, certificate, .txt

    def test_unregistered_client_rejected(self, tmp_path):
        share = tmp_path / "share" / "DP"
        _make_dated_docx(share / "DP-2026" / "KHAZAR" / "KF-DP-26-002.docx",
                         "10.01.2026")
        _make_dated_docx(share / "DP-2026" / "OTHER" / "XYZ-DP-26-002.docx",
                         "10.01.2026")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["KF-DP-26-002.docx"]
        assert r.rejected == 1

    def test_utility_subfolder_pruned(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2026" / "Templates" / "U-UT-26-900.docx",
                         "10.01.2026")
        _make_dated_docx(share / "UT-2026" / "UBOC" / "U-UT-26-901.docx",
                         "10.01.2026")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["U-UT-26-901.docx"]

    def test_legacy_doc_reported_not_staged(self, tmp_path):
        # A valid report name with the legacy .doc extension: python-docx cannot open
        # it, so it is logged for manual conversion rather than silently staged.
        share = tmp_path / "share" / "UT"
        _write(share / "UT-2026" / "SDLN" / "UT-SDLN-26-045.doc", b"legacy-word")
        settings = _settings(tmp_path, [share])
        lines: list[str] = []

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None, log=lines.append)

        assert r.staged == []
        assert any("legacy .doc not stageable" in line for line in lines)


class TestDateCutoffSources:
    def test_filename_date_gates_a_pdf(self, tmp_path):
        # A .pdf has no readable internal Date: — its filename tail is the only signal.
        share = tmp_path / "share" / "DP"
        _write(share / "DP-2025" / "SDLN" / "DP-SDLN-25-004(F1798)_12.09.2025.pdf",
               b"%PDF-legacy")
        _write(share / "DP-2025" / "SDLN" / "DP-SDLN-25-060(F1798)_16.12.2025.pdf",
               b"%PDF-current")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["DP-SDLN-25-060(F1798)_16.12.2025.pdf"]
        assert r.legacy_skipped == 1

    def test_undated_2025_pdf_dropped(self, tmp_path):
        # No filename date, no readable body: a 2025 scan-only report is legacy.
        share = tmp_path / "share" / "DP"
        _write(share / "DP-2025" / "KHAZAR" / "KF-DP-25-700.pdf", b"%PDF-scan")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == []
        assert r.legacy_skipped == 1

    def test_undated_2026_pdf_kept(self, tmp_path):
        # 2026 is wholly after the cutoff — staged without needing any date at all.
        share = tmp_path / "share" / "DP"
        _write(share / "DP-2026" / "KHAZAR" / "KF-DP-26-700.pdf", b"%PDF-scan")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["KF-DP-26-700.pdf"]
        assert r.legacy_skipped == 0

    def test_filename_date_beats_internal_date(self, tmp_path):
        # The SDLN filename tail is authoritative: a pre-cutoff tail drops the file
        # even though the body carries a post-cutoff date.
        share = tmp_path / "share" / "DP"
        _make_dated_docx(share / "DP-2025" / "SDLN" / "DP-SDLN-25-004_12.09.2025.docx",
                         "20.12.2025")
        settings = _settings(tmp_path, [share])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == []
        assert r.legacy_skipped == 1

    def test_twin_pdf_dropped_across_roots(self, tmp_path):
        # The .docx twin may sit under a different configured root than the .pdf.
        share_a = tmp_path / "share" / "DP"
        share_b = tmp_path / "share" / "DP-mirror"
        _make_dated_docx(share_a / "DP-2026" / "GEMZA" / "G-DP-26-001.docx",
                         "05.01.2026")
        _write(share_b / "DP-2026" / "GEMZA" / "G-DP-26-001.pdf", b"%PDF-twin")
        settings = _settings(tmp_path, [share_a, share_b])

        r = ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                          timeout=None)

        assert r.staged == ["G-DP-26-001.docx"]
        assert r.twin_skipped == 1


# --- live progress streaming -------------------------------------------------------------


class TestLiveProgressStream:
    def test_progress_lines_are_numbered_and_labelled(self, tmp_path):
        share = tmp_path / "share" / "MPI"
        _make_dated_docx(share / "MPI-2026" / "GEMZA" / "G-MT-26-001.docx",
                         "05.01.2026")
        _write(share / "MPI-2026" / "GEMZA" / "G-MT-26-001.pdf", b"%PDF-twin")
        settings = _settings(tmp_path, [share])
        lines: list[str] = []

        ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                      timeout=None, log=lines.append)

        progress = [line for line in lines if line.startswith("[LAN Sync ")]
        assert len(progress) == 2                       # one line per candidate
        assert progress[0].startswith("[LAN Sync 1/2] ")
        assert progress[1].startswith("[LAN Sync 2/2] ")
        joined = "\n".join(progress)
        assert "Staging G-MT-26-001.docx..." in joined
        assert "Skipped twin PDF G-MT-26-001.pdf..." in joined

    def test_scan_line_emitted_per_root(self, tmp_path):
        share = tmp_path / "share" / "UT"
        _make_dated_docx(share / "UT-2026" / "UBOC" / "U-UT-26-002.docx", "05.01.2026")
        settings = _settings(tmp_path, [share])
        lines: list[str] = []

        ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                      timeout=None, log=lines.append)

        assert any(line.startswith("[LAN Sync] Scanning ") for line in lines)

    def test_stream_writes_and_flushes_stdout(self, capsys):
        # The default log is _stream: straight to sys.stdout, flushed per line, so a
        # long walk never looks frozen behind a block-buffered pipe.
        ni._stream("[LAN Sync 1/1] Staging G-MT-26-001.docx...")
        assert capsys.readouterr().out == "[LAN Sync 1/1] Staging G-MT-26-001.docx...\n"

    def test_legacy_skip_line_names_the_reason(self, tmp_path):
        share = tmp_path / "share" / "DP"
        _write(share / "DP-2025" / "SDLN" / "DP-SDLN-25-004_12.09.2025.pdf",
               b"%PDF-old")
        settings = _settings(tmp_path, [share])
        lines: list[str] = []

        ni.sync_since(settings, CUTOFF, manifest_path=_manifest_path(tmp_path),
                      timeout=None, log=lines.append)

        assert any("Skipped legacy" in line and "2025-09-12 < 2025-12-01" in line
                   for line in lines)
