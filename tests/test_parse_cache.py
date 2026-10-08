"""Pin the persistent parse cache.

The property that matters: a second run over an unchanged corpus parses NOTHING and
still returns the same ``(cache, errors, textless)`` a cold run does — while an edited
file re-parses and a transient error is never remembered. These are the guarantees that
let ``main.py`` re-parse only new/changed files each month without ever handing the
aggregator a stale table.

Real ``.docx`` fixtures are used (not stubbed tables) so the cache is exercised through
the genuine ``docx_reader.read_many`` path, process pool and all. "Did it re-parse?" is
asserted by spying on ``docx_reader.read_many`` — the one symbol ``read_many_cached``
delegates a miss batch to, in-process, so the spy is reliable on spawn platforms too.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from monthly_report import docx_reader, parse_cache


def _make_docx(path: Path, *, cell_text: str | None = "hello") -> Path:
    """Write a minimal real .docx. ``cell_text=None`` makes a textless (image-like) doc."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    if cell_text is not None:
        table.cell(0, 0).text = cell_text
        table.cell(0, 1).text = "world"
    doc.save(str(path))
    return path


@pytest.fixture
def corpus(tmp_path: Path) -> tuple[Path, list[Path]]:
    """A (cache_dir, [docx...]) fixture with two text reports and one textless one."""
    src = tmp_path / "input"
    src.mkdir()
    files = [
        _make_docx(src / "a.docx", cell_text="alpha"),
        _make_docx(src / "b.docx", cell_text="beta"),
        _make_docx(src / "c.docx", cell_text=None),   # textless
    ]
    return tmp_path / ".parse_cache", files


def test_contract_parity_on_cold_cache(corpus):
    """read_many_cached returns the same shapes as read_many on a fresh store."""
    cache_dir, files = corpus
    plain_cache, plain_err, plain_textless = docx_reader.read_many(files)
    c_cache, c_err, c_textless = parse_cache.read_many_cached(files, cache_dir)

    assert set(c_cache) == set(plain_cache)
    assert c_cache.keys() == plain_cache.keys()
    assert set(c_err) == set(plain_err) == set()
    assert set(c_textless) == set(plain_textless)
    # tables round-trip through JSON identically to a live parse
    for path in plain_cache:
        assert c_cache[path] == plain_cache[path]


def test_hit_skips_parse(corpus, monkeypatch):
    """A second run over unchanged files delegates nothing to read_many."""
    cache_dir, files = corpus
    # Cold run populates the store.
    parse_cache.read_many_cached(files, cache_dir)

    calls = {"n": 0}
    real = docx_reader.read_many

    def _spy(paths, **kw):
        calls["n"] += 1
        return real(paths, **kw)

    monkeypatch.setattr(parse_cache.docx_reader, "read_many", _spy)

    cache, errors, textless = parse_cache.read_many_cached(files, cache_dir)
    assert calls["n"] == 0                    # every file served from disk
    assert set(cache) == set(files)
    assert parse_cache.read_many_cached.last_stats["hits"] == len(files)
    assert parse_cache.read_many_cached.last_stats["misses"] == 0
    # textless flag survives the round-trip from the store
    assert [p.name for p in textless] == ["c.docx"]


def test_edit_forces_reparse(corpus, monkeypatch):
    """Changing a file's bytes changes its hash, so it re-parses (a fresh miss)."""
    cache_dir, files = corpus
    parse_cache.read_many_cached(files, cache_dir)

    _make_docx(files[0], cell_text="ALPHA-EDITED")   # same path, new bytes

    seen = {}
    real = docx_reader.read_many

    def _spy(paths, **kw):
        seen["paths"] = list(paths)
        return real(paths, **kw)

    monkeypatch.setattr(parse_cache.docx_reader, "read_many", _spy)
    cache, _err, _tl = parse_cache.read_many_cached(files, cache_dir)

    assert seen["paths"] == [files[0]]                # only the edited file re-parsed
    assert parse_cache.read_many_cached.last_stats["misses"] == 1
    assert any("ALPHA-EDITED" in row for tbl in cache[files[0]] for grid in tbl for row in grid)


def test_errors_are_not_cached(tmp_path):
    """A corrupt file lands in errors, writes no entry, and re-tries cleanly later."""
    cache_dir = tmp_path / ".parse_cache"
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a real zip/docx")

    cache, errors, _tl = parse_cache.read_many_cached([bad], cache_dir)
    assert bad in errors and bad not in cache
    # No entry file was written for the unreadable file.
    digest = parse_cache.content_hash(bad)
    assert not (cache_dir / f"{digest}.json").exists()

    # Replaced with a valid document under the same path -> parses, no stale hit.
    _make_docx(bad, cell_text="recovered")
    cache2, errors2, _ = parse_cache.read_many_cached([bad], cache_dir)
    assert bad in cache2 and not errors2


def test_gc_prunes_only_unreferenced_entries(corpus):
    """gc keeps entries for files still present, drops the rest."""
    cache_dir, files = corpus
    parse_cache.read_many_cached(files, cache_dir)
    store = parse_cache.read_many_cached.last_stats["store"]
    live = parse_cache.read_many_cached.last_stats["live_digests"]

    assert len(list(cache_dir.glob("*.json"))) == len(files)
    # Nothing stale yet: a GC with the live set removes nothing.
    assert store.gc(live) == 0
    assert len(list(cache_dir.glob("*.json"))) == len(files)
    # An empty live set means every file has left input/: all entries retire.
    assert store.gc(set()) == len(files)
    assert list(cache_dir.glob("*.json")) == []


def test_schema_bump_invalidates_entry(corpus, monkeypatch):
    """An entry written under a different schema is a miss, never a misread."""
    cache_dir, files = corpus
    one = [files[0]]
    parse_cache.read_many_cached(one, cache_dir)
    digest = parse_cache.content_hash(files[0])
    assert (cache_dir / f"{digest}.json").exists()

    monkeypatch.setattr(parse_cache, "CACHE_SCHEMA", parse_cache.CACHE_SCHEMA + 1)
    assert parse_cache.ParseCache(cache_dir).get(digest) is None
