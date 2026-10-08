"""Pin the reader dispatch and what a disagreement does to the cache (ADR-003 M4).

Three properties:

1. **A .docx is unaffected.** The dispatcher must be invisible on the native path
   — same tables, same flag, no cross-check, no PyMuPDF.
2. **A .pdf carries its disagreements out.** They travel on
   ``read_many.last_disagreements`` / ``read_many_cached.last_disagreements``, not
   on the returned triple, so no existing caller changes.
3. **A blocking disagreement is never cached.** Otherwise the cross-check is a
   one-run event: quarantine it, put it back unchanged, and the second run serves
   cached tables and compiles the verdict nobody ever settled.
"""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from docx import Document

from monthly_report import docx_reader, parse_cache, pdf_crosscheck


def _make_docx(path: Path) -> Path:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Report No:"
    table.cell(0, 1).text = "MT-U-25-042"
    table.cell(1, 0).text = "Client:"
    table.cell(1, 1).text = "UBOC"
    doc.save(str(path))
    return path


def _make_pdf(path: Path) -> Path:
    """A ruled two-row table — the primary reader works off ruling lines, and every
    PDF in the corpus is Word output, so a borderless page would not be a fixture of
    anything real."""
    doc = fitz.open()
    page = doc.new_page()
    rows, cols = (90, 110, 130), (40, 200, 300, 400)
    for y in rows:
        page.draw_line((cols[0], y), (cols[-1], y))
    for x in cols:
        page.draw_line((x, rows[0]), (x, rows[-1]))
    for (x, y), text in zip(
            [(cols[0], rows[0]), (cols[1], rows[0]), (cols[2], rows[0]),
             (cols[0], rows[1]), (cols[1], rows[1]), (cols[2], rows[1])],
            ["Report No: RT-U-25-176", "Weld No", "Result", "W1", "6", "Accept"]):
        page.insert_text((x + 3, y + 13), text, fontsize=9)
    doc.save(str(path))
    doc.close()
    return path


class TestNativeDocxUntouched:
    def test_dispatch_returns_exactly_what_read_document_returns(self, tmp_path: Path) -> None:
        path = _make_docx(tmp_path / "r.docx")
        tables, has_text = docx_reader.read_document(path)
        name, dispatched, flag, found = docx_reader._read_one(str(path))
        assert (name, dispatched, flag, found) == (str(path), tables, has_text, [])

    def test_a_docx_batch_reports_no_disagreements(self, tmp_path: Path) -> None:
        paths = [_make_docx(tmp_path / f"r{i}.docx") for i in range(2)]
        cache, errors, textless = docx_reader.read_many(paths, max_workers=1)
        assert len(cache) == 2 and errors == {} and textless == []
        assert docx_reader.read_many.last_disagreements == []


class TestPdfDispatch:
    def test_a_pdf_is_read_directly_and_carries_a_disagreement_list(self, tmp_path: Path) -> None:
        path = _make_pdf(tmp_path / "r.pdf")
        name, tables, has_text, found = docx_reader._read_one(str(path))
        assert name == str(path) and has_text is True
        assert isinstance(found, list)
        assert any("RT-U-25-176" in cell
                   for table in tables for row in table for cell in row)

    def test_an_unreadable_pdf_becomes_an_error_not_a_crash(self, tmp_path: Path) -> None:
        broken = tmp_path / "broken.pdf"
        broken.write_text("not a pdf at all", encoding="utf-8")
        _name, result, has_text, found = docx_reader._read_one(str(broken))
        assert isinstance(result, str) and has_text is False and found == []


class TestBlockingIsNeverCached:
    """The property that keeps a disagreement from being a one-run event."""

    @pytest.fixture
    def blocking(self, tmp_path: Path, monkeypatch):
        path = _make_pdf(tmp_path / "r.pdf")

        def _stub(path_str: str):
            found = [pdf_crosscheck.Disagreement(
                report_no="RT-U-25-176", field="result", primary="Fit",
                crosscheck="Unfit", path=path_str, blocking=True)]
            return path_str, [[["cell"]]], True, found

        monkeypatch.setattr(docx_reader, "_read_one", _stub)
        return path

    def test_nothing_is_stored_for_a_blocking_pdf(self, blocking: Path, tmp_path: Path) -> None:
        cache_dir = tmp_path / ".cache"
        parse_cache.read_many_cached([blocking], cache_dir, max_workers=1)
        assert list(cache_dir.glob("*.json")) == []

    def test_the_disagreement_reaches_the_caller(self, blocking: Path, tmp_path: Path) -> None:
        parse_cache.read_many_cached([blocking], tmp_path / ".cache", max_workers=1)
        found = parse_cache.read_many_cached.last_disagreements
        assert [(d.field, d.blocking) for d in found] == [("result", True)]

    def test_a_second_run_re_parses_it(self, blocking: Path, tmp_path: Path) -> None:
        cache_dir = tmp_path / ".cache"
        parse_cache.read_many_cached([blocking], cache_dir, max_workers=1)
        parse_cache.read_many_cached([blocking], cache_dir, max_workers=1)
        stats = parse_cache.read_many_cached.last_stats
        assert stats["hits"] == 0 and stats["misses"] == 1

    def test_a_clean_pdf_is_cached_normally(self, tmp_path: Path, monkeypatch) -> None:
        path = _make_pdf(tmp_path / "clean.pdf")
        monkeypatch.setattr(docx_reader, "_read_one",
                            lambda s: (s, [[["cell"]]], True, []))
        cache_dir = tmp_path / ".cache"
        parse_cache.read_many_cached([path], cache_dir, max_workers=1)
        assert len(list(cache_dir.glob("*.json"))) == 1
        parse_cache.read_many_cached([path], cache_dir, max_workers=1)
        assert parse_cache.read_many_cached.last_stats["hits"] == 1
